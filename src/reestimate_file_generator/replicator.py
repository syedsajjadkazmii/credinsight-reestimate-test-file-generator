from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.workbook.workbook import Workbook

from reestimate_file_generator.config import FILE_CONFIGS, PORTFOLIO_FILENAME, FileConfig
from reestimate_file_generator.id_generator import collect_existing_ids, generate_new_ids


@dataclass
class LoanAssignment:
    output_id: str
    source_id: str


@dataclass
class GenerationResult:
    output_dir: Path
    assignments: list[LoanAssignment]
    file_stats: dict[str, dict[str, int]]


def read_portfolio_loan_ids(input_dir: Path) -> list[str]:
    path = input_dir / PORTFOLIO_FILENAME
    if not path.exists():
        raise FileNotFoundError(f"Portfolio file not found: {path}")

    cfg = next(c for c in FILE_CONFIGS if c.filename == PORTFOLIO_FILENAME)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[cfg.sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        raise ValueError("Portfolio file is empty")

    loan_ids: list[str] = []
    for row in rows[1:]:
        if row and row[cfg.loan_col_idx]:
            loan_ids.append(str(row[cfg.loan_col_idx]))
    if not loan_ids:
        raise ValueError("No loan IDs found in portfolio file")
    return loan_ids


def build_assignments(
    base_loan_ids: list[str], count: int, input_dir: Path
) -> list[LoanAssignment]:
    if count < 1:
        raise ValueError("Count must be at least 1")
    if count > len(base_loan_ids):
        clone_count = count - len(base_loan_ids)
        reserved = collect_existing_ids(input_dir)
        new_ids = generate_new_ids(clone_count, reserved)
        assignments = [
            LoanAssignment(output_id=lid, source_id=lid) for lid in base_loan_ids
        ]
        for i, new_id in enumerate(new_ids):
            source_id = base_loan_ids[i % len(base_loan_ids)]
            assignments.append(LoanAssignment(output_id=new_id, source_id=source_id))
        return assignments

    selected = base_loan_ids[:count]
    return [LoanAssignment(output_id=lid, source_id=lid) for lid in selected]


def _read_file_rows(path: Path, cfg: FileConfig) -> tuple[tuple, list[tuple]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[cfg.sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return (), []

    header = rows[0]
    data_rows = [row for row in rows[1:] if row and row[cfg.loan_col_idx]]
    return header, data_rows


def _build_loan_row_map(
    data_rows: list[tuple], loan_col_idx: int, allowed_ids: set[str]
) -> dict[str, list[tuple]]:
    loan_map: dict[str, list[tuple]] = {}
    for row in data_rows:
        loan_id = str(row[loan_col_idx])
        if loan_id not in allowed_ids:
            continue
        loan_map.setdefault(loan_id, []).append(row)
    return loan_map


def _clone_row(row: tuple, loan_col_idx: int, new_loan_id: str) -> tuple:
    row_list = list(row)
    row_list[loan_col_idx] = new_loan_id
    return tuple(row_list)


def _write_output_file(
    cfg: FileConfig,
    header: tuple,
    output_rows: list[tuple],
    output_path: Path,
) -> None:
    wb: Workbook = openpyxl.Workbook()
    ws = wb.active
    ws.title = cfg.sheet_name

    for col_idx, value in enumerate(header, start=1):
        ws.cell(row=1, column=col_idx, value=value)

    for row_idx, row in enumerate(output_rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    wb.save(output_path)
    wb.close()


def generate_files(
    input_dir: Path,
    output_dir: Path,
    count: int,
) -> GenerationResult:
    input_dir = input_dir.resolve()
    output_dir = output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    base_loan_ids = read_portfolio_loan_ids(input_dir)
    allowed_ids = set(base_loan_ids)
    assignments = build_assignments(base_loan_ids, count, input_dir)

    file_stats: dict[str, dict[str, int]] = {}

    for cfg in FILE_CONFIGS:
        input_path = input_dir / cfg.filename
        if not input_path.exists():
            raise FileNotFoundError(f"Input file not found: {input_path}")

        header, data_rows = _read_file_rows(input_path, cfg)
        loan_map = _build_loan_row_map(data_rows, cfg.loan_col_idx, allowed_ids)

        output_rows: list[tuple] = []
        for assignment in assignments:
            source_rows = loan_map.get(assignment.source_id, [])
            if not source_rows:
                continue
            for row in source_rows:
                if assignment.output_id == assignment.source_id:
                    output_rows.append(row)
                else:
                    output_rows.append(
                        _clone_row(row, cfg.loan_col_idx, assignment.output_id)
                    )

        output_path = output_dir / cfg.filename
        _write_output_file(cfg, header, output_rows, output_path)

        unique_loans = {str(row[cfg.loan_col_idx]) for row in output_rows}
        file_stats[cfg.filename] = {
            "rows": len(output_rows),
            "unique_loans": len(unique_loans),
        }

    return GenerationResult(
        output_dir=output_dir,
        assignments=assignments,
        file_stats=file_stats,
    )


def validate_result(result: GenerationResult) -> list[str]:
    """Return a list of validation errors (empty if all checks pass)."""
    errors: list[str] = []
    expected_loans = {a.output_id for a in result.assignments}
    expected_count = len(expected_loans)

    portfolio_stats = result.file_stats.get(PORTFOLIO_FILENAME, {})
    if portfolio_stats.get("unique_loans") != expected_count:
        errors.append(
            f"Portfolio has {portfolio_stats.get('unique_loans')} loans, "
            f"expected {expected_count}"
        )

    for filename, stats in result.file_stats.items():
        if stats["unique_loans"] != expected_count:
            errors.append(
                f"{filename}: {stats['unique_loans']} unique loans, "
                f"expected {expected_count}"
            )

    return errors
